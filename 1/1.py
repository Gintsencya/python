from docx import Document

doc = Document("1.docx")

table = doc.add_table(3,5)
for i in range(0,5):
    a = table.cell(1,i)
    a.text ='sssaaa'
doc.save("1.docx")

print(doc)